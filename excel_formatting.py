"""Excel formatting helpers to match the original RA Rate card appearance."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RA_SHEET_NAME = "Rate card"
TITLE_ROW = 1
SUBTITLE_ROW = 2
COST_NAME_ROW = 4
VALIDITY_ROW = 5
RATE_BY_ROW = 6
HEADER_ROW = 8
DATA_START_ROW = 9

TITLE_FONT = Font(name="Arial", size=18, color="0063A3")
SUBTITLE_FONT = Font(name="Arial", size=15, color="0063A3")
HEADER_FONT = Font(name="Arial", size=10)
HEADER_FONT_MUTED = Font(name="Arial", size=10, color="808080")
DATA_FONT = Font(name="Arial", size=10)
META_FONT = Font(name="Arial", size=10)

HEADER_FILL = PatternFill("solid", fgColor="F1F1F6")
INVOICE_TYPE_FILL = PatternFill("solid", fgColor="D9D9D9")

LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

DEFAULT_COLUMN_WIDTH = 13.0
RATE_NUMBER_FORMAT = "General"

MUTED_HEADER_NAMES = {
    "tab",
    "origin country",
    "origin state",
    "origin city",
    "uplift airport",
    "rate type",
    "destination country",
    "destination state",
    "destination",
    "country of clearance",
    "service",
    "lane type",
}

DISPLAY_HEADER_OVERRIDES = {
    "Origin City_2": "Origin City",
    "Destination_2": "Destination",
    "Lane Type_2": "Lane Type",
    "Service_2": "Service",
}

RATE_BY_TEXT = "Rate by: Weight/chargeable kg\r\nRegular rule"


def _cell_text(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def _display_header(column_name: str) -> str:
    return DISPLAY_HEADER_OVERRIDES.get(column_name, column_name)


def _header_font(column_name: str) -> Font:
    if _display_header(column_name).casefold() in MUTED_HEADER_NAMES:
        return HEADER_FONT_MUTED
    return HEADER_FONT


def _read_title_rows(source_ra_file: Path) -> tuple[str, str]:
    workbook = load_workbook(source_ra_file, data_only=True, read_only=True)
    try:
        worksheet = workbook[RA_SHEET_NAME]
        title = _cell_text(worksheet.cell(TITLE_ROW, 1).value) or "Rate Card"
        subtitle = _cell_text(worksheet.cell(SUBTITLE_ROW, 1).value)
        return title, subtitle
    finally:
        workbook.close()


def _split_rate_card_columns(df: pd.DataFrame) -> tuple[list[str], str, str, str]:
    columns = list(df.columns)
    currency_idx = next(
        idx for idx, column in enumerate(columns) if _cell_text(column).casefold() == "currency"
    )
    shipment_columns = columns[:currency_idx]
    currency_column = columns[currency_idx]
    surcharge_columns = columns[currency_idx + 1 :]
    if len(surcharge_columns) != 1:
        raise ValueError("Expected exactly one fuel surcharge column after Currency")
    surcharge_column = surcharge_columns[0]
    if not surcharge_column.startswith("Fuel Surcharge ("):
        raise ValueError(f"Unexpected fuel surcharge column name: {surcharge_column}")
    return shipment_columns, currency_column, surcharge_column, surcharge_column


def _style_header_cell(cell, column_name: str) -> None:
    cell.font = _header_font(column_name)
    cell.fill = HEADER_FILL
    cell.alignment = LEFT_ALIGN


def _style_meta_cell(cell) -> None:
    cell.font = META_FONT
    cell.alignment = LEFT_ALIGN


def _style_data_cell(cell, *, column_name: str, value: object) -> None:
    cell.font = DATA_FONT
    cell.alignment = LEFT_ALIGN
    if column_name.casefold() == "invoice type":
        cell.fill = INVOICE_TYPE_FILL
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.number_format = RATE_NUMBER_FORMAT


def save_formatted_rate_card(
    df: pd.DataFrame,
    output_path: Path,
    *,
    source_ra_file: Path,
    sheet_name: str = RA_SHEET_NAME,
) -> Path:
    shipment_columns, currency_column, surcharge_column, surcharge_title = _split_rate_card_columns(df)
    title, subtitle = _read_title_rows(source_ra_file)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]

    currency_col_idx = len(shipment_columns) + 1
    surcharge_col_idx = currency_col_idx + 1
    last_col_idx = surcharge_col_idx

    worksheet.cell(TITLE_ROW, 1, title).font = TITLE_FONT
    if subtitle:
        worksheet.cell(SUBTITLE_ROW, 1, subtitle).font = SUBTITLE_FONT

    worksheet.merge_cells(
        start_row=COST_NAME_ROW,
        start_column=currency_col_idx,
        end_row=COST_NAME_ROW,
        end_column=surcharge_col_idx,
    )
    cost_name_cell = worksheet.cell(COST_NAME_ROW, currency_col_idx, surcharge_title)
    _style_meta_cell(cost_name_cell)

    worksheet.merge_cells(
        start_row=RATE_BY_ROW,
        start_column=currency_col_idx,
        end_row=RATE_BY_ROW,
        end_column=surcharge_col_idx,
    )
    rate_by_cell = worksheet.cell(RATE_BY_ROW, currency_col_idx, RATE_BY_TEXT)
    _style_meta_cell(rate_by_cell)

    for col_idx, column_name in enumerate(shipment_columns, start=1):
        header_cell = worksheet.cell(HEADER_ROW, col_idx, _display_header(column_name))
        _style_header_cell(header_cell, column_name)

    currency_header = worksheet.cell(HEADER_ROW, currency_col_idx, "Currency")
    _style_header_cell(currency_header, "Currency")

    surcharge_header = worksheet.cell(HEADER_ROW, surcharge_col_idx, "p/unit")
    _style_header_cell(surcharge_header, "p/unit")

    for row_offset, (_, row) in enumerate(df.iterrows()):
        excel_row = DATA_START_ROW + row_offset

        for col_idx, column_name in enumerate(shipment_columns, start=1):
            value = row[column_name]
            if _cell_text(value):
                cell = worksheet.cell(excel_row, col_idx, value)
            else:
                cell = worksheet.cell(excel_row, col_idx)
            _style_data_cell(cell, column_name=column_name, value=value)

        currency_value = row[currency_column]
        currency_cell = worksheet.cell(excel_row, currency_col_idx)
        if _cell_text(currency_value):
            currency_cell.value = currency_value
        _style_data_cell(currency_cell, column_name="Currency", value=currency_value)

        surcharge_value = row[surcharge_column]
        surcharge_cell = worksheet.cell(excel_row, surcharge_col_idx)
        if pd.notna(surcharge_value):
            surcharge_cell.value = surcharge_value
        _style_data_cell(surcharge_cell, column_name="p/unit", value=surcharge_value)

    for col_idx in range(1, last_col_idx + 1):
        worksheet.column_dimensions[get_column_letter(col_idx)].width = DEFAULT_COLUMN_WIDTH

    worksheet.freeze_panes = worksheet.cell(DATA_START_ROW, 1)
    worksheet.sheet_view.showGridLines = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
